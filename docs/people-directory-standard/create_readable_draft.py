from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
ROOT=Path(__file__).resolve().parent
D=Document();sec=D.sections[0]
sec.page_width=Inches(8.5);sec.page_height=Inches(11)
sec.top_margin=sec.bottom_margin=Inches(.65);sec.left_margin=sec.right_margin=Inches(.75)
for name in ['Normal','Body Text','List Bullet']:
 st=D.styles[name];st.font.name='Arial';st.font.size=Pt(11);st.paragraph_format.space_after=Pt(8);st.paragraph_format.line_spacing=1.08
for name,size in [('Title',26),('Heading 1',21),('Heading 2',13)]:
 st=D.styles[name];st.font.name='Arial';st.font.size=Pt(size);st.font.color.rgb=RGBColor.from_string('0C2340');st.paragraph_format.space_before=Pt(12);st.paragraph_format.space_after=Pt(7)
header=sec.header.paragraphs[0];header.text='INSTITUTE WEBSITES  /  PROPOSED PEOPLE DIRECTORY STANDARD';header.runs[0].font.name='Arial';header.runs[0].font.size=Pt(9)
footer=sec.footer.paragraphs[0];footer.text='Draft for feedback  •  23 October 2026  •  v0.2                    ';footer.runs[0].font.size=Pt(9)
footer.add_run('Page ');field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
D.core_properties.title='People directories: a shared guide for institute communicators'
D.core_properties.subject='Proposed standards — revised for readability from the supplied PDF'
D.core_properties.author='Keough School of Global Affairs'
base='https://prugamba-nd.github.io/people-hub-mockup/'
def p(text,style=None):return D.add_paragraph(text,style)
def h(text):D.add_heading(text,2)
def bullet(text):p(text,'List Bullet')
def link(par,label,url):
 rel=par.part.relate_to(url,RT.HYPERLINK,is_external=True);node=OxmlElement('w:hyperlink');node.set(qn('r:id'),rel)
 run=OxmlElement('w:r');pr=OxmlElement('w:rPr');color=OxmlElement('w:color');color.set(qn('w:val'),'0563C1');pr.append(color);u=OxmlElement('w:u');u.set(qn('w:val'),'single');pr.append(u);run.append(pr);t=OxmlElement('w:t');t.text=label;run.append(t);node.append(run);par._p.append(node)
def example(label,path):link(D.add_paragraph(),label,base+path)
def table(head,rows,widths):
 t=D.add_table(rows=1,cols=len(head));t.style='Table Grid';t.autofit=False
 for col,width in zip(t.columns,widths):col.width=Inches(width)
 for i,text in enumerate(head):t.rows[0].cells[i].text=text
 for row in rows:
  cells=t.add_row().cells
  for i,text in enumerate(row):cells[i].text=text
 for n,row in enumerate(t.rows):
  trpr=row._tr.get_or_add_trPr();trpr.append(OxmlElement('w:cantSplit'))
  if n==0:trpr.append(OxmlElement('w:tblHeader'))
  for i,cell in enumerate(row.cells):
   cell.width=Inches(widths[i])
   for par in cell.paragraphs:
    par.paragraph_format.space_before=Pt(5);par.paragraph_format.space_after=Pt(5)
    for run in par.runs:run.font.name='Arial';run.font.size=Pt(10.5)
   if n==0:
    shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'0C2340');cell._tc.get_or_add_tcPr().append(shade)
    for run in cell.paragraphs[0].runs:run.bold=True;run.font.color.rgb=RGBColor(255,255,255)
 return t
D.add_heading('People directories',0)
p('A shared guide for institute communicators','Subtitle')
p('These proposed standards give visitors a consistent way to find people across institute websites, whatever theme the site uses. They apply to institute sites, not the Keough School website.')
D.add_heading('1. Set up the directory',1)
h('Start with People in the main navigation')
p('Clicking People opens a page titled People Directory. No sidebar category is selected on this landing page.')
bullet('Go straight to the profiles. Do not add a large introductory image or inset hero.')
bullet('Use four profile columns on desktop, with fewer columns on smaller screens.')
bullet('Show leadership first, then everyone else alphabetically by surname. List each person once.')
bullet('Use five sidebar groups: Leadership, Faculty, Staff, Scholars & Fellows, and Students.')
h('Add controls as the directory grows')
p('Use the same thresholds for the main directory and each category or subcategory listing.')
t=table(['Profiles','How to display them','Example'],[
('1–24','Show everyone on one page.',''),
('25–40','Show 24 per page, with pagination.',''),
('41 or more','Show 24 per page, with pagination, search, and relevant filters.','')],[.8,4.6,1.6])
for i,(label,path) in enumerate([('24 profiles','screens/24/'),('40 profiles','screens/40/'),('60 profiles','screens/60/')],1):link(t.rows[i].cells[2].paragraphs[0],label,base+path)
p('Category and subcategory listings are alphabetical by surname. Only the main directory puts leadership first.')
par=p('See the directory in each theme: ')
for i,(label,path) in enumerate([('NDT4',''),('NDT3','ndt3/'),('Theme 2.4','ndt24/')]):
 if i:par.add_run('  ·  ')
 link(par,label,base+path)
D.add_page_break()
D.add_heading('2. Organize the groups',1)
p('Each sidebar category opens its own page. That page shows the category’s default group and reveals its additional groups in the sidebar.')
h('Keep category pages simple')
bullet('Use one page title, such as Faculty or Advisory. Do not add introductory text or extra section headings above or between profiles. Keep names and role labels on the cards.')
bullet('Show the default group directly. For example, Faculty shows core affiliated faculty; do not add a separate “Core Affiliated Faculty” heading or child link.')
bullet('Reveal child groups only for the selected category. Keep other categories collapsed and mark the selected page.')
table(['Category','Who appears on the page','Additional groups in the sidebar'],[
('Leadership','Institute leadership','Advisory'),
('Faculty','Core affiliated faculty','Faculty Fellows; Associated Faculty; Emeriti'),
('Staff','Institute staff','None'),
('Scholars & Fellows','Visiting scholars and fellows','Postdoctoral Researchers; Affiliates; Past Scholars & Fellows; PhDs on the market'),
('Students','Undergraduate and graduate students, excluding PhD students','PhD')],[1.35,2.2,3.45])
h('Give navigation, search, and filters different jobs')
bullet('Sidebar: move to a different group, such as Faculty or Staff.')
bullet('Search: find someone within the current listing. Hub search covers everyone; Faculty search covers faculty.')
bullet('Filters: narrow the current listing by useful attributes, such as discipline or region. Do not duplicate the sidebar groups.')
example('Try category search and filters','compare/category/')
D.add_page_break()
D.add_heading('3. Choose where each profile leads',1)
p('First, identify who maintains the full biography. Not every directory entry needs a biography page on the institute website.')
h('Full biography maintained by the Keough School')
p('Link the directory entry directly to the school’s profile. Do not create an intermediate institute biography page. The external-link arrow shows that the visitor is leaving the institute site.')
example('Example: directory card linking to the school','profile-examples/external-profile/')
h('Full biography maintained by the institute')
p('Open a detailed local biography page containing the person’s full profile information.')
example('Example: full institute biography','profile-examples/full-bio/')
h('Full biography maintained elsewhere')
p('When institute-specific context is useful, open an abridged local biography. Describe the person’s work or relationship with the institute and include a button linking to the full biography outside the Keough School website.')
example('Example: abridged biography with a full-bio button','profile-examples/abridged-bio/')
h('Historical affiliation, without a maintained biography')
p('For groups such as former visiting professors, use names only or images and names. These entries acknowledge a past affiliation; they do not promise a current biography.')
bullet('Do not link entries by default or create a local biography page.')
bullet('Keep names and historical affiliations accurate. Do not copy biographies your team cannot keep current.')
p('The profile-card thresholds on page 1 do not set a separate rule for these simple historical lists.')
par=p('Examples: ');link(par,'Names only',base+'profile-examples/names-only/');par.add_run('  ·  ');link(par,'Images and names',base+'profile-examples/images-and-names/')
h('Before publishing')
p('Confirm the person’s group, who maintains the full biography, and which of the four options above applies.')
D.save(ROOT/'Institute-People-Directory-Standard-Readable-Draft.docx')
print('Created revised Word draft.')
