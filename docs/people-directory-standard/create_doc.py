from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.opc.constants import RELATIONSHIP_TYPE as RT
ROOT=Path(__file__).resolve().parent
D=Document();section=D.sections[0]
section.top_margin=Inches(.65);section.bottom_margin=Inches(.6);section.left_margin=Inches(.75);section.right_margin=Inches(.75)
section.page_width=Inches(8.5);section.page_height=Inches(11)
styles=D.styles
for name in ['Normal','Body Text']:
 styles[name].font.name='Calibri';styles[name].font.size=Pt(11)
 styles[name].paragraph_format.space_after=Pt(7)
 styles[name].paragraph_format.line_spacing=1.08
for name,size in [('Title',27),('Heading 1',19),('Heading 2',13)]:
 styles[name].font.name='Calibri';styles[name].font.size=Pt(size);styles[name].font.color.rgb=RGBColor.from_string('0C2340')
 styles[name].paragraph_format.space_before=Pt(12);styles[name].paragraph_format.space_after=Pt(7)
styles['Caption'].font.size=Pt(9);styles['Caption'].font.color.rgb=RGBColor.from_string('555555')
header=section.header.paragraphs[0];header.text='INSTITUTE WEBSITES  /  SHARED PEOPLE DIRECTORY STANDARD';header.runs[0].font.size=Pt(9);header.runs[0].font.color.rgb=RGBColor.from_string('536173')
foot=section.footer.paragraphs[0];foot.text='DRAFT FOR FEEDBACK  •  23 October 2026  •  v0.1';foot.runs[0].font.size=Pt(9)
foot.add_run(' '*8+'Page ')
f=OxmlElement('w:fldSimple');f.set(qn('w:instr'),'PAGE');foot._p.append(f)
D.core_properties.title='People pages: proposed shared standard for institute websites'
D.core_properties.subject='Draft for institute communicator review'
D.core_properties.author='Keough School of Global Affairs'
D.core_properties.keywords='People directory, institute communicators, proposed standard'
def p(t='',style=None):return D.add_paragraph(t,style)
def h(t):D.add_heading(t,2)
def bullet(t):p(t,'List Bullet')
def link(label,url):
 par=D.add_paragraph();rel=par.part.relate_to(url,RT.HYPERLINK,is_external=True);hl=OxmlElement('w:hyperlink');hl.set(qn('r:id'),rel);r=OxmlElement('w:r');prop=OxmlElement('w:rPr');c=OxmlElement('w:color');c.set(qn('w:val'),'0563C1');prop.append(c);u=OxmlElement('w:u');u.set(qn('w:val'),'single');prop.append(u);r.append(prop);t=OxmlElement('w:t');t.text=label;r.append(t);hl.append(r);par._p.append(hl)
def table(headers,rows,widths=None):
 t=D.add_table(rows=1, cols=len(headers));t.style='Table Grid'
 for i,x in enumerate(headers):t.rows[0].cells[i].text=x
 for row in rows:
  cells=t.add_row().cells
  for i,x in enumerate(row):cells[i].text=x
 for ri,row in enumerate(t.rows):
  pr=row._tr.get_or_add_trPr();cant=OxmlElement('w:cantSplit');pr.append(cant)
  if ri==0:
   repeat=OxmlElement('w:tblHeader');pr.append(repeat)
  for ci,cell in enumerate(row.cells):
   if widths:cell.width=Inches(widths[ci])
   for par in cell.paragraphs:
    par.paragraph_format.space_after=Pt(5);par.paragraph_format.space_before=Pt(4)
    for run in par.runs:run.font.size=Pt(10)
   if ri==0:
    shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'0C2340');cell._tc.get_or_add_tcPr().append(shade)
    for run in cell.paragraphs[0].runs:run.bold=True;run.font.color.rgb=RGBColor(255,255,255)
 return t
D.add_heading('People pages',0)
p('Proposed shared standard for institute websites','Subtitle')
p('For institute communicators • Draft for feedback • 23 October 2026')
p('This proposal establishes a consistent way to find people and maintain their profiles across institute websites, regardless of the website theme. It applies to institute sites; it does not prescribe the structure of the Keough School website.')
h('1. Make People a top-level navigation item')
p('Every institute website uses People in its main navigation. Clicking it opens the People Directory, with no category selected in the sidebar.')
h('2. Open directly on the directory')
bullet('Use People Directory as the page title. Do not place an inset hero or introductory photo above the directory.')
bullet('Show profile cards in four columns on desktop, adapting to fewer columns on smaller screens.')
bullet('List leadership first, followed by everyone else in alphabetical order by surname. Show each person once in the hub listing.')
bullet('Provide these five sidebar categories: Leadership; Faculty; Staff; Scholars & Fellows; Students.')
h('3. Scale listings by the number of profiles')
p('Apply these thresholds to the full directory and to category or subcategory listings:')
table(['Profiles in the listing','What visitors see'],[('1–24','All profiles on one page.'),('25–40','Pagination, with 24 profiles per page.'),('41 or more','Pagination, with 24 profiles per page, plus search and filters appropriate to the listing.')],[1.5,5.5])
p('Category and subcategory listings are alphabetical by surname. Leadership-first ordering applies only to the main People Directory.')
p('Search on the hub covers the directory. Search within a category covers that category; filters should help narrow its profiles, rather than duplicate the sidebar groups.')
D.add_page_break()
D.add_heading('Category pages and the sidebar',1)
p('A category page answers “Who belongs to this group?” It opens a dedicated listing rather than expanding a long list of every subcategory on the hub.')
h('4. Use one page title')
bullet('Use only the category name as the page heading—for example, Faculty or Advisory.')
bullet('Do not add a lede or extra section headings above or between the profiles. Profile names and role labels remain part of the cards.')
bullet('When a category is selected, reveal its lower-level categories in the sidebar. Keep the other categories collapsed and indicate the selected page.')
h('5. Let the parent category show its default group')
p('The parent page shows the group visitors would naturally expect there. Do not repeat that default group as a child link or an extra heading. For example, Faculty shows core affiliated faculty without a separate “Core Affiliated Faculty” heading or sidebar item.')
table(['Sidebar category','Profiles shown on the parent page','Children revealed when selected'],[
('Leadership','Institute leadership','Advisory'),
('Faculty','Core affiliated faculty','Faculty Fellows; Associated Faculty; Emeriti'),
('Staff','Institute staff','None in the proposed structure'),
('Scholars & Fellows','Visiting scholars and fellows together','Postdoctoral Researchers; Affiliates; Past Scholars & Fellows'),
('Students','Undergraduate and graduate students, excluding PhD students','PhD')],[1.35,2.55,3.1])
h('Sidebar, search, and filters have different jobs')
p('The sidebar moves visitors between groups. Search finds a person or matching information within the current listing. Filters narrow that listing using useful attributes, such as discipline or region of expertise, when the profile information supports them.')
p('Example: a visitor opens Faculty, searches for migration, then narrows the results to a discipline. They should not need to navigate through additional sidebar levels for every research topic.')
D.add_page_break()
D.add_heading('Profile destinations and historical listings',1)
h('6. Follow where the full biography is maintained')
p('Institute directories include people whose primary affiliation is with the institute and people whose full profile is maintained by another entity. A directory listing does not always need a local biography page.')
table(['Where the full biography is maintained','What clicking the directory entry opens'],[
('Keough School website','The full profile on the school website, directly. Do not create an intermediate institute biography page.'),
('The institute website','A detailed institute biography page containing the person’s full profile information.'),
('A host outside the Keough School website, with institute-specific context to provide','An abridged institute biography: a short account of the person’s work or relationship specific to the institute, plus a link to the full profile on its host site.')],[2.5,4.5])
p('Some directory cards will therefore take visitors directly off the institute website. Local biography pages have two forms: a full profile maintained by the institute, or an abridged institute-specific profile linking to the full biography elsewhere.')
h('7. Use simple listings when current biographies cannot be maintained')
p('Historical groups—such as visiting professors who have left—may use a plain listing instead of clickable profile cards. These entries record a past relationship; they do not imply that the institute maintains a current biography.')
bullet('Use a name alone, or an image and name.')
bullet('Entries are not linked by default and do not require a local biography page.')
bullet('Do not reproduce biographies that the institute cannot keep current. Basic accuracy of names and historical affiliation still matters.')
p('This is a distinct category-page format. It is appropriate when the purpose is to acknowledge past affiliation rather than provide a current profile directory. The profile-card pagination thresholds above do not establish a separate threshold for these simple historical lists.')
h('Before publishing a profile entry')
p('Confirm the appropriate group, where the full biography is maintained, and whether the entry should open a full institute profile, an abridged institute profile, a Keough School profile, or remain a historical name-only entry.')
D.add_page_break()
D.add_heading('Examples and feedback',1)
p('Use this illustration and the linked mockups to discuss the proposed experience. The written rules in this draft are the proposal; the exploratory mockups are not all approved patterns.')
D.add_picture(str(ROOT/'assets/hub-illustration.png'),width=Inches(6.8))
p('Illustrative hub with fictional profiles: People Directory, no active sidebar category, leadership first, and four columns. This view was adapted from the prototype for this document.','Caption')
base='https://prugamba-nd.github.io/people-hub-mockup/'
link('Try the large directory: search, filters, and pagination',base+'compare/directory/')
link('Try Faculty: category-specific search, filters, and child navigation',base+'compare/category/')
link('Compare theme examples: NDT4',base)
link('Compare theme examples: NDT3',base+'ndt3/')
link('Compare theme examples: Theme 2.4',base+'ndt24/')
p('Example limitations: the large mockup uses fictional records and alphabetical ordering; existing live hubs still use the earlier “People” title. The document illustration and written rules specify the proposed title and leadership-first order. Profile-destination and historical-list rules are described here but are not demonstrated fully in these mockups.','Caption')
h('Please add comments in Word')
p('Use Review → New Comment beside the rule you are responding to. Please identify the institute or example that informs your feedback.')
for t in ['Does the category structure fit your institute? Identify a group or parent-page default that needs discussion.', 'Can your team identify and maintain the correct profile destination for each person?', 'Are the listing thresholds workable, and what information would make useful filters for your categories?']:
 bullet(t)
D.save(ROOT/'Institute-People-Directory-Standard-Draft.docx')
print(ROOT/'Institute-People-Directory-Standard-Draft.docx')
