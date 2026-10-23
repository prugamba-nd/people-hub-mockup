from pathlib import Path
import json
from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR
from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.lib.colors import HexColor
ROOT=Path(__file__).resolve().parent;ASSETS=ROOT/'assets'
META=json.loads((ASSETS/'capture-metadata.json').read_text())
META['hub']['boxes']['first']['w']=META['hub']['boxes']['title']['w']
NAVY='0C2340';GOLD='AE9142';GRAY='52616D';LINE='D5DBDF';WHITE='FFFFFF'
pdfmetrics.registerFont(TTFont('Arial','/System/Library/Fonts/Supplemental/Arial.ttf'))
pdfmetrics.registerFont(TTFont('ArialBold','/System/Library/Fonts/Supplemental/Arial Bold.ttf'))
pres=Presentation();pres.slide_width=Inches(16);pres.slide_height=Inches(9)
pres.core_properties.title='People directories: a visual guide for institute communicators'
pres.core_properties.subject='Proposed standards — screenshot-led draft for feedback'
pres.core_properties.author='Keough School of Global Affairs'
pdf=canvas.Canvas(str(ROOT/'People-Directory-Visual-Guide.pdf'),pagesize=(16*72,9*72))
pdf.setTitle('People directories: a visual guide for institute communicators')
base='https://prugamba-nd.github.io/people-hub-mockup/'
slide=None;number=0

def rect(x,y,w,h,fill=None,stroke=None,width=1):
 sh=slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h))
 if fill:sh.fill.solid();sh.fill.fore_color.rgb=RGBColor.from_string(fill)
 else:sh.fill.background()
 if stroke:sh.line.color.rgb=RGBColor.from_string(stroke);sh.line.width=Pt(width)
 else:sh.line.fill.background()
 pdf.saveState()
 if fill:pdf.setFillColor(HexColor('#'+fill))
 if stroke:pdf.setStrokeColor(HexColor('#'+stroke));pdf.setLineWidth(width)
 pdf.rect(x*72,(9-y-h)*72,w*72,h*72,fill=bool(fill),stroke=bool(stroke));pdf.restoreState()
 return sh

def text(x,y,w,h,txt,size=16,bold=False,color=NAVY,url=None):
 sh=slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));tf=sh.text_frame;tf.clear();tf.word_wrap=True
 tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0;tf.vertical_anchor=MSO_ANCHOR.TOP
 for i,line in enumerate(txt.split('\n')):
  p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.space_after=Pt(0);p.line_spacing=1.15
  run=p.add_run();run.text=line;run.font.name='Arial';run.font.size=Pt(size);run.font.bold=bold;run.font.color.rgb=RGBColor.from_string(color)
  if url:run.hyperlink.address=url
 pdf.saveState();pdf.setFillColor(HexColor('#'+color));pdf.setFont('ArialBold' if bold else 'Arial',size)
 for i,line in enumerate(txt.split('\n')):pdf.drawString(x*72,(9-y)*72-size-i*size*1.15,line)
 if url:pdf.linkURL(url,(x*72,(9-y-h)*72,(x+w)*72,(9-y)*72),relative=0)
 pdf.restoreState();return sh

def badge(x,y,n,r=.17):
 sh=slide.shapes.add_shape(MSO_SHAPE.OVAL,Inches(x-r),Inches(y-r),Inches(2*r),Inches(2*r));sh.fill.solid();sh.fill.fore_color.rgb=RGBColor.from_string(NAVY);sh.line.fill.background()
 tf=sh.text_frame;tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0;tf.vertical_anchor=MSO_ANCHOR.MIDDLE
 p=tf.paragraphs[0];p.alignment=2;run=p.add_run();run.text=str(n);run.font.name='Arial';run.font.size=Pt(13);run.font.bold=True;run.font.color.rgb=RGBColor(255,255,255)
 pdf.saveState();pdf.setFillColor(HexColor('#'+NAVY));pdf.circle(x*72,(9-y)*72,r*72,fill=1,stroke=0);pdf.setFillColor(HexColor('#'+WHITE));pdf.setFont('ArialBold',13);pdf.drawCentredString(x*72,(9-y)*72-4.3,str(n));pdf.restoreState()

def image(name,x,y,w,h=None,url=None):
 path=ASSETS/(name+'.png');iw,ih=Image.open(path).size
 if h is None:h=w*ih/iw
 pic=slide.shapes.add_picture(str(path),Inches(x),Inches(y),width=Inches(w),height=Inches(h))
 if url:pic.click_action.hyperlink.address=url
 pdf.drawImage(str(path),x*72,(9-y-h)*72,w*72,h*72)
 if url:pdf.linkURL(url,(x*72,(9-y-h)*72,(x+w)*72,(9-y)*72),relative=0)
 rect(x,y,w,h,stroke=LINE,width=.6)
 return (x,y,w,h)

def image_region(name,x,y,w,crop,url=None):
 path=ASSETS/(name+'.png');iw,ih=Image.open(path).size
 left,top,cw,ch=crop;h=w*ch/cw
 pic=slide.shapes.add_picture(str(path),Inches(x),Inches(y),width=Inches(w),height=Inches(h))
 pic.crop_left=left/iw;pic.crop_right=(iw-left-cw)/iw;pic.crop_top=top/ih;pic.crop_bottom=(ih-top-ch)/ih
 if url:pic.click_action.hyperlink.address=url
 pdf.saveState();clip=pdf.beginPath();clip.rect(x*72,(9-y-h)*72,w*72,h*72);pdf.clipPath(clip,stroke=0,fill=0)
 scale=w/cw;fullh=ih*scale
 pdf.drawImage(str(path),(x-left*scale)*72,(9-y+top*scale-fullh)*72,iw*scale*72,fullh*72)
 pdf.restoreState();rect(x,y,w,h,stroke=LINE,width=.6)
 if url:pdf.linkURL(url,(x*72,(9-y-h)*72,(x+w)*72,(9-y)*72),relative=0)
 return x,y,w,h

def highlight(name,key,placement,n,maxheight=None):
 m=META[name];b=m['boxes'][key];c=m['clip'];x,y,w,h=placement
 bx=max(0,b['x']-c['x']);by=max(0,b['y']-c['y']);bw=min(b['w'],c['width']-bx);bh=min(b['h'],c['height']-by)
 if maxheight:bh=min(bh,maxheight)
 sx=x+bx*w/c['width'];sy=y+by*h/c['height'];sw=bw*w/c['width'];sh=bh*h/c['height']
 rect(sx-.035,sy-.035,sw+.07,sh+.07,stroke=GOLD,width=2)
 badge(sx,sy,n)

def callout(n,y,title,body=''):
 badge(11.65,y+.14,n);text(12,y,3.5,.55,title,18,True)
 if body:text(12,y+.64,3.45,1.0,body,14,color=GRAY)

def new(title,kicker,notes,source):
 global slide,number
 if slide is not None:pdf.showPage()
 number+=1;slide=pres.slides.add_slide(pres.slide_layouts[6]);rect(0,0,16,.055,fill=GOLD)
 text(.5,.2,15,.2,'PROPOSED STANDARD  /  INSTITUTE PEOPLE PAGES',9,bold=True,color=GRAY)
 text(.5,.55,15,.6,title,30,True)
 text(.5,8.72,10,.18,'DRAFT FOR FEEDBACK  ·  '+kicker,9,color=GRAY)
 text(12.4,8.72,2.6,.18,'Open the live example',9,color=GRAY,url=source)
 text(15.25,8.72,.5,.18,f'{number} / 7',9,color=GRAY)
 slide.notes_slide.notes_text_frame.text=notes+'\n\nSource / example: '+source+'\nIllustrative institute mockups. The underlying website components use NDT4. Screenshots may show fictional people; public profile examples cite their sources on the linked pages.'

new('People opens the directory','Start here','Use People in top-level navigation. The landing title is People Directory. No category is selected on the hub. The five categories are Leadership, Faculty, Staff, Scholars & Fellows, and Students. Four columns on desktop adapt to smaller screens. Leadership comes first; everyone else follows alphabetically by surname. Show each person once. No inset hero or introductory photo. Screenshot uses the 40-profile NDT4 example.',base+'screens/40/')
pos=image('hub',.5,1.25,10.65,url=base+'screens/40/')
highlight('hub','nav',pos,1);highlight('hub','title',pos,2);highlight('hub','sidebar',pos,3,maxheight=330);highlight('hub','first',pos,4,maxheight=410)
callout(1,1.7,'People in the\nmain navigation')
callout(2,3.1,'One clear title','People Directory')
callout(3,4.4,'Five sidebar groups','None selected on the hub.')
callout(4,5.9,'Four columns','Leadership first.\nThen surname A–Z.')

new('The number of profiles determines the controls','Directory + category listings','Use these thresholds for the main directory, categories, and subcategories. Up to 24: all entries on one page, no pagination. 25–40: 24 per page, native pagination below the grid. More than 40: same pagination plus scoped search and relevant filters. 40 means two pages (24 and 16); 60 means three pages (24, 24, and 12). Screenshots intentionally show top-of-page and bottom-of-list excerpts rather than shrinking a whole long page. The 24-record illustration is rendered using the native NDT4 40-profile example reduced locally to 24 records for this deck. No live page was changed.',base+'screens/')
for i,(n,summary) in enumerate([(24,'All on one page'),(40,'24 per page'),(60,'24 per page + search / filters')]):
 x=.5+i*5.17
 text(x,1.4,4.7,.45,{24:'Up to 24 profiles',40:'25–40 profiles',60:'More than 40 profiles'}[n],26,True)
 text(x,1.98,4.8,.35,summary,14,color=GRAY)
 image('size-'+str(n)+'-top',x,2.6,4.65,url=base+f'screens/{n}/')
 if n>24:
  text(x,6.05,4.6,.25,'BOTTOM OF THE LIST',10,bold=True,color=GRAY)
  image_region('size-'+str(n)+'-bottom',x,6.48,4.65,(250,15,500,90),url=base+f'screens/{n}/')
  rect(x-.04,6.43,4.73,.94,stroke=GOLD,width=2)
  text(x,7.6,4.8,.4,'Numbered pagination below the grid',12,True)
 else:
  text(x,6.42,4.7,.5,'No pagination.\nNo search or filters.',18,True)
 if n==60:rect(x-.04,3.35,4.73,1.42,stroke=GOLD,width=2)

new('A category opens its default group','Browse by group','Each category opens a dedicated page, with one title and no lede or additional section headings. The default group is implicit: Faculty displays core affiliated faculty. Additional groups appear only beneath the selected parent. This capture removes the old extra Core Affiliated Faculty heading and sorts profiles by surname locally to illustrate the proposed standard; the live category example may still contain it. Parent defaults: Leadership=institute leadership; Faculty=core affiliated faculty; Staff=institute staff; Scholars & Fellows=visiting scholars and fellows; Students=undergraduate and graduate, excluding PhD. The reviewed draft also lists PhDs on the market under Scholars & Fellows; retain for institute feedback.',base+'faculty/')
pos=image('category',.5,1.25,10.65,url=base+'faculty/')
highlight('category','title',pos,1);highlight('category','children',pos,2);highlight('category','first',pos,3,maxheight=430)
callout(1,1.7,'Just “Faculty”','No lede.\nNo repeated group heading.')
callout(2,3.55,'Reveal its subgroups','Other categories\nstay collapsed.')
callout(3,5.45,'Show core affiliates','No separate “Core”\npage or heading needed.')

new('A school-hosted profile opens directly','Profile destination 1','When the full biography is maintained on the Keough School website, the directory card links directly to that profile. Do not create an intermediate institute biography page. NDT4 supplies the external-link arrow automatically for external card URLs. The public school example is Scott Appleby; the card’s link and photo source are recorded in the prototype.',base+'profile-examples/external-profile/')
pos=image('external',.5,1.45,10.65,url=base+'profile-examples/external-profile/')
highlight('external','link',pos,1)
callout(1,2.0,'The arrow signals\nan offsite destination','Clicking the name opens\nthe school’s full profile.')
text(12,5.1,3.5,1.4,'No intermediate\ninstitute bio page.',18,True)

new('A full biography lives on the institute site','Profile destination 2','Use the native biography layout shown on the Ann Firth page: name and role at the top, biography at left, portrait and contact details at right. The institute maintains the full biography. This example uses fictional Maren Okafor and reserved example.org contact details. Layout reference: https://president.nd.edu/about/leadership-and-staff/ann-firth/',base+'profile-examples/full-bio/')
pos=image('full',.5,1.35,8.5,url=base+'profile-examples/full-bio/')
highlight('full','title',pos,1,maxheight=160);highlight('full','text',pos,2,maxheight=350);highlight('full','portrait',pos,3,maxheight=520)
callout(1,1.7,'Name + role')
callout(2,3.25,'Full biography','Maintained by the institute.')
callout(3,5.15,'Photo + contact details','Keep office and email\nwith the portrait.')

new('An abridged biography adds institute context','Profile destination 3','Use the same native biography layout. Keep the local text short and specific to the institute relationship. Retain office and email beneath the portrait. Link the button to the full biography maintained outside the Keough School website. The example uses Daniel Philpott’s public department profile. The complete biography, publications and contact information are maintained by the department; the local text explains the institute connection.',base+'profile-examples/abridged-bio/')
pos=image('abridged',.5,1.45,10.25,url=base+'profile-examples/abridged-bio/')
highlight('abridged','text',pos,1,maxheight=245);highlight('abridged','portrait',pos,2,maxheight=520);highlight('abridged','button',pos,3)
callout(1,1.75,'Short, institute-specific\nbiography')
callout(2,3.6,'Photo + office + email')
callout(3,5.2,'Button to the full bio','Hosted outside the\nKeough School website.')

new('Past affiliation can be a simple list','Historical listings','Use these for historical groups whose current biographies the institute cannot maintain, such as former visiting professors. Names only or images and names; entries are not linked by default and do not require a local biography. Do not reproduce biographies your team cannot keep current. Check basic accuracy of names and historical affiliation. No separate pagination threshold for simple historical lists is established by this proposal. Both examples use fictional names. Image-and-name example: '+base+'profile-examples/images-and-names/',base+'profile-examples/names-only/')
text(.55,1.35,6.8,.45,'Names only',22,True);text(8.3,1.35,7,.45,'Images + names',22,True)
image('names',.55,2.05,7.1,url=base+'profile-examples/names-only/')
image('images',8.3,2.05,7.1,url=base+'profile-examples/images-and-names/')
rect(.55,7.95,14.85,.52,fill=NAVY)
text(.85,8.06,14,.28,'No profile links.  No biography to keep current.  Names and affiliations still need to be accurate.',16,True,WHITE)

pdf.save();pres.save(ROOT/'People-Directory-Visual-Guide.pptx')
print('Created 7-slide editable PowerPoint and matching PDF.')
